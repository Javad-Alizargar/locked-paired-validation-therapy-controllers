"""Build PeerJ-formatted tables (markdown + .docx)."""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Inches

HERE = Path(__file__).resolve().parent
RES = HERE / "results"
OUT = HERE / "tables"
OUT.mkdir(exist_ok=True)

CTRL_ORDER = ["Smoothed predictive", "Thermal-floor predictive", "Constrained optimized", "Reduced-order MPC"]
STRAT_ORDER = ["all", "low", "transitional", "high"]
STRAT_NAME = {"all": "All", "low": "Predicted-low", "transitional": "Predicted-intermediate", "high": "Predicted-high"}


def f3(x: float, dp: int = 3) -> str:
    return f"{x:+.{dp}f}"


def f0(x: float, dp: int = 3) -> str:
    return f"{x:.{dp}f}"


def ci(r: pd.Series, col: str, dp: int = 3) -> str:
    lo = f"{r[col + '_ci_low']:.{dp}f}"
    hi = f"{r[col + '_ci_high']:.{dp}f}"
    return f"{f3(r[col], dp)} ({lo} to {hi})"


def pfmt(p: float) -> str:
    if pd.isna(p):
        return "n/a"
    return "<0.001" if p < 0.001 else f"{p:.3f}"


def docx_table(title: str, note: str, headers: list[str], rows: list[list[str]], path: Path, widths=None):
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(12)
    for s in doc.sections:
        s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Inches(1)
        s.page_width = Inches(8.5); s.page_height = Inches(11)

    h = doc.add_paragraph()
    run = h.add_run(title)
    run.bold = True
    run.font.size = Pt(11)

    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, htxt in enumerate(headers):
        hdr[i].text = ""
        par = hdr[i].paragraphs[0]
        r = par.add_run(htxt)
        r.bold = True
        r.font.size = Pt(8)
        r.font.name = "Times New Roman"
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            par = cells[i].paragraphs[0]
            r = par.add_run(val)
            r.font.size = Pt(8)
            r.font.name = "Times New Roman"
    if widths:
        for i, w in enumerate(widths):
            for row in t.rows:
                row.cells[i].width = Inches(w)

    n = doc.add_paragraph()
    r = n.add_run(note)
    r.font.size = Pt(9)
    r.italic = True
    doc.save(path)
    print(f"  {path.name}")


def md_table(title: str, note: str, headers: list[str], rows: list[list[str]], path: Path):
    lines = [f"**{title}**", "", "| " + " | ".join(headers) + " |", "| " + " | ".join(["---"] * len(headers)) + " |"]
    for r in rows:
        lines.append("| " + " | ".join(r) + " |")
    lines += ["", f"*{note}*", ""]
    path.write_text("\n".join(lines), encoding="utf-8")
    print(f"  {path.name}")


def main() -> None:
    p = pd.read_csv(RES / "primary_endpoints.csv")
    par = pd.read_csv(RES / "pareto_summary.csv").set_index("controller")
    p["corder"] = p.controller.map({c: i for i, c in enumerate(CTRL_ORDER)})
    p["sorder"] = p.stratum.map({s: i for i, s in enumerate(STRAT_ORDER)})

    # ---------------- Table 1: success endpoints
    h1 = ["Controller", "Pairs", "Baseline success", "Controller success",
          "Change in success (95% CI)", "p", "MDE (80% power)", "Rescued / lost"]
    r1 = []
    for _, r in p[p.stratum == "all"].sort_values("corder").iterrows():
        r1.append([
            r.controller, f"{int(r.n_pairs)}", f0(r.baseline_success_rate),
            f0(r.controller_success_rate), ci(r, "delta_success"),
            pfmt(r.delta_success_p), f0(r.delta_success_mde80),
            f"{int(r.rescued)} / {int(r.lost)}",
        ])
    note1 = ("All 4,608 paired trials per controller (384 clusters; 12 scenarios per cluster). "
             "Change is the paired risk difference against the same full-feedback baseline, with a 95% "
             "cluster-bootstrap interval (10,000 replicates) and a cluster-level sign-flip p-value "
             "(20,000 permutations). MDE is the minimum detectable effect at 80% power given the observed "
             "cluster-robust standard error. These all-stratum contrasts are descriptive; the locked "
             "multiplicity-controlled family is the 12 controller × stratum tests reported in Table 3.")
    docx_table("Table 1. Success endpoints across the full locked validation atlas.", note1, h1, r1,
               OUT / "Table 1.docx", widths=[1.4, 0.5, 0.65, 0.7, 1.5, 0.5, 0.7, 0.65])
    md_table("Table 1. Success endpoints across the full locked validation atlas.", note1, h1, r1, OUT / "Table 1.md")

    # ---------------- Table 2: mechanism + energy
    h2 = ["Controller", "Thermal failure", "ROS failure", "Overheating",
          "Delivered energy (J cm⁻²)", "Pareto status"]
    r2 = []
    for _, r in p[p.stratum == "all"].sort_values("corder").iterrows():
        r2.append([
            r.controller,
            ci(r, "delta_thermal_failure"), ci(r, "delta_ros_failure"),
            ci(r, "delta_overheat_failure"), ci(r, "delta_energy_j_cm2", 1),
            "Non-dominated" if par.loc[r.controller, "pareto_optimal"] else "Dominated",
        ])
    note2 = ("Values are paired changes in the endpoint rate (or in delivered energy) against the full-feedback "
             "baseline, with 95% cluster-bootstrap intervals (10,000 replicates). Negative failure changes indicate "
             "fewer failures. Hypoxic failure did not change for any controller (all deltas 0.000, 95% CI 0.000 to "
             "0.000) and is omitted. Pareto status is computed over the five objectives shown in Table 1 and this "
             "table; all four controllers are mutually non-dominated.")
    docx_table("Table 2. Failure-mode and energy endpoints.", note2, h2, r2,
               OUT / "Table 2.docx", widths=[1.5, 1.35, 1.3, 1.3, 1.55, 0.9])
    md_table("Table 2. Failure-mode and energy endpoints.", note2, h2, r2, OUT / "Table 2.md")

    # ---------------- Table 3: stratified success
    h3 = ["Controller", "Stratum", "Pairs", "Baseline success", "Change in success (95% CI)",
          "p", "Holm-adjusted p", "MDE (80% power)"]
    r3 = []
    for _, r in p.sort_values(["sorder", "corder"]).iterrows():
        r3.append([
            r.controller, STRAT_NAME[r.stratum], f"{int(r.n_pairs)}",
            f0(r.baseline_success_rate), ci(r, "delta_success"),
            pfmt(r.delta_success_p), pfmt(r.delta_success_p_holm) if r.stratum != "all" else "n/a",
            f0(r.delta_success_mde80),
        ])
    note3 = ("Strata are terciles of a frozen logistic scoring rule for predicted baseline success. The rule was "
             "fitted on independent seeds and applied unchanged here, so stratum membership depends only on "
             "pre-treatment design quantities and never on the controller outcome. The 12 controller × stratum "
             "tests form the primary family for Holm adjustment. Exact p-values are reported to three decimal "
             "places; the minimum attainable p-value with 20,000 permutations is 0.00005.")
    docx_table("Table 3. Success change by predicted-baseline-success stratum.", note3, h3, r3,
               OUT / "Table 3.docx", widths=[1.4, 1.15, 0.5, 0.7, 1.5, 0.5, 0.75, 0.7])
    md_table("Table 3. Success change by predicted-baseline-success stratum.", note3, h3, r3, OUT / "Table 3.md")

    # ---------------- Supplementary tables
    ends = pd.read_csv(RES / "endpoint_definitions.csv")
    hs = ["Criterion", "Rule", "Meaning"]
    rs = [[r.criterion, r.rule, r.meaning] for _, r in ends.iterrows()]
    notes = ("Endpoint definitions were frozen before any controller outcome was examined. warm_fraction is the "
             "fraction of the protocol window with 43 ≤ T ≤ 50 °C; overheat_fraction is the fraction with T > 55 °C; "
             "ros_threshold_fraction is the fraction with ROS dose at or above the singlet-oxygen threshold; "
             "hypoxic_failure_fraction is the fraction with O2 < 5 µM and ROS below threshold.")
    docx_table("Table S2. Frozen endpoint definitions.", notes, hs, rs, OUT / "Supplemental Table S2.docx", widths=[1.3, 2.4, 3.0])
    md_table("Table S2. Frozen endpoint definitions.", notes, hs, rs, OUT / "Table S2.md")

    st = pd.read_csv(RES / "strata_audit.csv")
    hs2 = ["Scope", "Seed", "Model family", "n", "AUC", "Observed success", "Mean predicted"]
    rs2 = [[r.scope, r.seed, r.model_family, f"{int(r.n)}", f0(r.auc), f0(r.observed_success), f0(r.mean_predicted)]
           for _, r in st.iterrows()]
    notes2 = ("Discrimination of the frozen scoring rule. The locked validation rows are the analysis population; "
              "the external row is a fully held-out seed that was never used to fit or to tune the rule. "
              "Discrimination is modest, which is why strata are reported as predicted-baseline-success bands "
              "rather than as a validated controllability construct.")
    docx_table("Table S1. Discrimination of the frozen scoring rule.", notes2, hs2, rs2,
               OUT / "Supplemental Table S1.docx", widths=[1.5, 0.9, 1.1, 0.4, 0.6, 0.9, 0.9])
    md_table("Table S1. Discrimination of the frozen scoring rule.", notes2, hs2, rs2, OUT / "Table S1.md")

    ws = pd.read_csv(RES / "weight_sensitivity.csv")
    hs3 = ["Controller", "P(rank 1)", "P(top 2)", "P(rank last)", "Mean rank"]
    rs3 = [[r.controller, f0(r.p_rank1), f0(r.p_top2), f0(r.p_rank_last), f0(r.mean_rank, 2)]
           for _, r in ws.iterrows()]
    notes3 = ("200,000 draws from a uniform Dirichlet prior over the five objectives (success, thermal failure, "
              "ROS failure, overheating, energy), each min-max normalised across controllers before weighting. "
              "This replaces the single fixed-weight utility used in the earlier analysis, which is not "
              "reproducible as a preference statement.")
    docx_table("Table S3. Rank stability under random objective weights.", notes3, hs3, rs3,
               OUT / "Supplemental Table S3.docx", widths=[1.7, 0.8, 0.8, 0.9, 0.8])
    md_table("Table S3. Rank stability under random objective weights.", notes3, hs3, rs3, OUT / "Table S3.md")


if __name__ == "__main__":
    print("writing tables to", OUT)
    main()
