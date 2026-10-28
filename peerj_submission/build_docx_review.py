"""Build a reviewer convenience copy of the manuscript with figures and
tables embedded inline at their first citation.

PeerJ requires the *submission* manuscript to carry figures and tables as
separate files (section 4.6/4.8: "Do not embed figures in the manuscript,
unless submitting a LaTeX manuscript").  This build therefore produces a
SECOND file, `main_manuscript_review_copy.docx`, intended for upload as a
Review-only Supplemental File (section 4.9) so that reviewers can read the
article with every figure and table in place.

Nothing in this file changes the submission manuscript or its numbering.
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt

from build_docx import (
    HERE,
    SRC,
    TITLE,
    add_runs,
    body,
    bullet,
    configure,
    heading,
    render_title_page,
)

FIG_DIR = HERE.parent / "peerj_analysis" / "figures"
TABLE_MD = HERE.parent / "peerj_analysis" / "tables"
CAPTIONS = HERE / "figure_captions.md"
OUT = HERE / "main_manuscript_review_copy.docx"

FIG_CITE = re.compile(r"\bFig\.\s*(\d+)")
TAB_CITE = re.compile(r"\bTable\s*(S?\d+)\b")


# ------------------------------------------------------------------ captions
def load_captions() -> dict[str, str]:
    caps: dict[str, str] = {}
    for line in CAPTIONS.read_text(encoding="utf-8").split("\n"):
        m = re.match(r"\*\*(Figure|Table) (S?\d+)\. (.*)$", line.strip())
        if m:
            caps[f"{m.group(1)} {m.group(2)}"] = line.strip()
    return caps


# -------------------------------------------------------------------- tables
def load_table(num: str):
    """Return (title, header, rows, footnote) parsed from Table <num>.md."""
    path = TABLE_MD / f"Table {num}.md"
    lines = [l.rstrip() for l in path.read_text(encoding="utf-8").split("\n")]
    title = next(l for l in lines if l.startswith("**Table"))
    rows = [l for l in lines if l.startswith("|")]
    header = [c.strip() for c in rows[0].strip("|").split("|")]
    data = [[c.strip() for c in r.strip("|").split("|")] for r in rows[2:]]
    foot = next((l for l in lines if l.startswith("*") and not l.startswith("**")), "")
    return title, header, data, foot


def add_table(doc: Document, num: str) -> None:
    """Reuse the verified standalone table, including widths and paragraph spacing."""
    from copy import deepcopy
    from docx.oxml.ns import qn
    filename = f"Supplemental Table {num}.docx" if num.startswith("S") else f"Table {num}.docx"
    source = Document(TABLE_MD / filename)
    for element in source.element.body:
        if element.tag != qn("w:sectPr"):
            doc.element.body.insert(len(doc.element.body) - 1, deepcopy(element))
    doc.add_paragraph()


def add_figure(doc: Document, num: int, caption: str) -> None:
    path = FIG_DIR / f"Figure {num}.png"
    if not path.exists():
        raise FileNotFoundError(path)
    par = doc.add_paragraph()
    par.paragraph_format.alignment = 1  # centred image, left text elsewhere
    run = par.add_run()
    # 6.2 in keeps every figure inside the 6.5 in text column.
    run.add_picture(str(path), width=Inches(6.2))
    cap = doc.add_paragraph()
    add_runs(cap, caption)
    for r in cap.runs:
        r.font.size = Pt(10)
    doc.add_paragraph()


# --------------------------------------------------------------------- build
def build() -> None:
    lines = SRC.read_text(encoding="utf-8").split("\n")
    caps = load_captions()

    doc = Document()
    configure(doc)
    render_title_page(doc, TITLE)

    from build_docx import page_setup  # local import: needs a section
    from docx.enum.section import WD_SECTION

    section = doc.add_section(WD_SECTION.NEW_PAGE)
    page_setup(section)

    seen: set[str] = set()
    for raw in lines:
        line = raw.rstrip()
        if not line.strip():
            continue
        if line.startswith("# "):
            heading(doc, line[2:].strip(), 1)
        elif line.startswith("## "):
            heading(doc, line[3:].strip(), 2)
        elif line.startswith("### "):
            heading(doc, line[4:].strip(), 3)
        elif line.startswith("- "):
            bullet(doc, line[2:].strip())
        else:
            text = line.strip()
            body(doc, text)

            for m in FIG_CITE.finditer(text):
                key = f"Figure {m.group(1)}"
                if key in caps and key not in seen:
                    seen.add(key)
                    add_figure(doc, int(m.group(1)), caps[key])

            for m in TAB_CITE.finditer(text):
                num = m.group(1)
                key = f"Table {num}"
                if key not in seen and (TABLE_MD / f"Table {num}.md").exists():
                    seen.add(key)
                    add_table(doc, num)

    doc.save(OUT)
    print(f"wrote {OUT.name}  (embedded: {sorted(seen)})")


if __name__ == "__main__":
    build()
