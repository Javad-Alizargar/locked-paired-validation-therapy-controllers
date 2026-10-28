"""Build the PeerJ Computer Science submission .docx from the markdown source.

Formatting required by the journal:
  * US Letter, 1 inch margins
  * 12 pt Times New Roman, left justified (never full justified)
  * continuous line numbers
  * unnumbered sections

The title page is rendered first as the author cover page (unnumbered
section), followed by a page break and the manuscript body with continuous
line numbers starting from the first line of the manuscript.
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

HERE = Path(__file__).resolve().parent
SRC = HERE / "main_manuscript.md"
TITLE = HERE / "title_page.md"
OUT = HERE / "main_manuscript.docx"

TOKEN = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+`)")


def add_runs(par, text: str) -> None:
    for part in TOKEN.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            run = par.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run = par.add_run(part[1:-1])
            run.italic = True
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            par.add_run(part[1:-1])
        else:
            par.add_run(part)


def add_line_numbers(section) -> None:
    sectPr = section._sectPr
    existing = sectPr.find(qn("w:lnNumType"))
    if existing is not None:
        sectPr.remove(existing)
    ln = OxmlElement("w:lnNumType")
    ln.set(qn("w:countBy"), "1")
    ln.set(qn("w:restart"), "continuous")
    ln.set(qn("w:distance"), "360")
    cols = sectPr.find(qn("w:cols"))
    if cols is not None:
        cols.addprevious(ln)
    else:
        sectPr.append(ln)


def page_setup(section) -> None:
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)


def configure(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), "Times New Roman")
    pf = normal.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf.space_before = Pt(0)
    pf.space_after = Pt(6)
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE

    page_setup(doc.sections[0])


def heading(doc: Document, text: str, level: int) -> None:
    par = doc.add_paragraph()
    par.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    par.paragraph_format.keep_with_next = True
    if level == 1:
        par.paragraph_format.space_before = Pt(0)
        par.paragraph_format.space_after = Pt(18)
    elif level == 2:
        par.paragraph_format.space_before = Pt(18)
        par.paragraph_format.space_after = Pt(6)
    else:
        par.paragraph_format.space_before = Pt(12)
        par.paragraph_format.space_after = Pt(4)
    run = par.add_run(text)
    run.font.name = "Times New Roman"
    if level == 1:
        run.font.size = Pt(16)
        run.bold = True
    elif level == 2:
        run.font.size = Pt(12)
        run.bold = True
    else:
        run.font.size = Pt(12)
        run.bold = True
        run.italic = True


def body(doc: Document, text: str, hanging: bool = False) -> None:
    par = doc.add_paragraph()
    par.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if hanging:
        par.paragraph_format.left_indent = Inches(0.3)
        par.paragraph_format.first_line_indent = Inches(-0.3)
    add_runs(par, text)


def bullet(doc: Document, text: str) -> None:
    par = doc.add_paragraph(style="List Bullet")
    par.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_runs(par, text)


def tp_heading(doc: Document, text: str) -> None:
    par = doc.add_paragraph()
    pf = par.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf.space_before = Pt(10)
    pf.space_after = Pt(1)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    pf.keep_with_next = True
    run = par.add_run(text)
    run.bold = True
    run.font.size = Pt(12)


def tp_body(doc: Document, text: str, size: int = 12, bold: bool = False) -> None:
    par = doc.add_paragraph()
    pf = par.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf.space_before = Pt(0)
    pf.space_after = Pt(1)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    add_runs(par, text)
    for run in par.runs:
        run.font.size = Pt(size)
        if bold:
            run.bold = True


def tp_bullet(doc: Document, text: str) -> None:
    par = doc.add_paragraph(style="List Bullet")
    pf = par.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf.space_before = Pt(0)
    pf.space_after = Pt(1)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    add_runs(par, text)


def render_title_page(doc: Document, src: Path) -> None:
    """Render the title page as the author cover page (no line numbers)."""
    lines = src.read_text(encoding="utf-8").split("\n")
    section_label = ""
    for raw in lines:
        line = raw.rstrip()
        if not line.strip():
            continue
        if line.startswith("# ") and not line.startswith("## "):
            continue  # drop the "# Title Page" document heading
        if line.startswith("## "):
            section_label = line[3:].strip()
            tp_heading(doc, section_label)
            continue
        if line.startswith("- "):
            tp_bullet(doc, line[2:].strip())
            continue
        if section_label.lower() == "article title":
            tp_body(doc, line.strip(), size=14, bold=True)
        else:
            tp_body(doc, line.strip())


def build(src: Path, out: Path) -> None:
    lines = src.read_text(encoding="utf-8").split("\n")
    doc = Document()
    configure(doc)

    render_title_page(doc, TITLE)

    # Start the manuscript in its own section so that line numbering begins
    # with the first line of the manuscript rather than on the cover page.
    manuscript_section = doc.add_section(WD_SECTION.NEW_PAGE)
    page_setup(manuscript_section)
    add_line_numbers(manuscript_section)

    i = 0
    in_references = False
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue
        if line.startswith("# "):
            heading(doc, line[2:].strip(), 1)
        elif line.startswith("## "):
            title = line[3:].strip()
            in_references = title.lower() == "references"
            heading(doc, title, 2)
        elif line.startswith("### "):
            heading(doc, line[4:].strip(), 3)
        elif line.startswith("- "):
            bullet(doc, line[2:].strip())
        else:
            # gather a full paragraph (markdown paragraphs are single lines here)
            body(doc, line.strip(), hanging=in_references)
        i += 1

    doc.save(out)
    print(f"wrote {out.name}")


if __name__ == "__main__":
    build(SRC, OUT)
